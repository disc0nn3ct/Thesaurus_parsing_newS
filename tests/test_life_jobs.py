"""Receipt work produces private review drafts; only explicit confirmation spends money."""
import csv
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from uuid import uuid4

from openpyxl import load_workbook

from thesaurus import life_jobs as jobs
from thesaurus import life_ledger as ledger
from thesaurus import receipt_wines, wine_retail
from thesaurus.config import Settings
from thesaurus.delivery import deliver_pending
from thesaurus.exporting import safe_cell
from thesaurus.feature_access import change, key
from thesaurus.service import Worker
from thesaurus.store import Store, dumps, utcnow


class LifeJobTests(unittest.TestCase):
    def setUp(self):
        parent = Path(__file__).resolve().parents[1] / '.test-tmp'
        parent.mkdir(exist_ok=True)
        directory = tempfile.TemporaryDirectory(dir=parent)
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.store = Store(self.root / 'state.sqlite3')
        self.settings = Settings(root=self.root)
        self.store.add_user(1, 'admin')
        for actor in (2, 3):
            self.store.add_user(actor, subscriptions=['life', 'vivino'])
        self.store.add_user(4)
        wine_retail.configure_shop(self.store, 1, 'decanter', True)
        self.enterContext(patch('requests.sessions.Session.request', side_effect=AssertionError('No network in ledger jobs tests')))

    def receipt(self, actor=2):
        path = self.settings.data_dir / 'receipts' / str(actor) / 'synthetic.pdf'
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b'Synthetic local receipt, mocked recognition')
        return path

    def candidate(self):
        return {'status': 'candidate', 'date': '2026-01-10', 'merchant': 'Synthetic shop',
                'total': 250, 'currency': 'RUB', 'fingerprint': 'a' * 64,
                'items': [{'name': 'Wine Test 2020 750 ml', 'quantity': 2., 'unit_price': 100,
                           'amount': 200, 'is_wine': True, 'vintage': 2020, 'volume_ml': 750, 'winery': ''},
                          {'name': 'Milk', 'quantity': 1., 'unit_price': 50, 'amount': 50,
                           'is_wine': False, 'vintage': None, 'volume_ml': None, 'winery': ''}],
                'warnings': ['Проверьте сумму'], 'requires_confirmation': True, 'raw_text': 'PRIVATE FULL OCR'}

    def payload(self, kind='life_receipt', **changes):
        return {'kind': kind}, {'actor_id': 2, 'owner_id': 2, 'report_id': uuid4().hex,
                               'request_id': uuid4().hex, **changes}

    def run_job(self, job, payload, **kwargs):
        return jobs.run_job(self.store, self.settings, job, payload, **kwargs)

    def report(self, payload):
        return self.store.one('SELECT * FROM reports WHERE id=?', (payload['report_id'],))

    def document(self, payload):
        return Path(json.loads(self.report(payload)['parts'])[0]['path'])

    def assert_no_financial_changes(self):
        self.assertFalse(self.store.query('SELECT * FROM life_ledger_entries'))
        self.assertFalse(self.store.query('SELECT * FROM life_receipt_imports'))
        self.assertFalse(self.store.query('SELECT * FROM wine_lots'))

    def seed_entries(self, count=1, actor=2, date='2026-01-10', **changes):
        with self.store.connection() as conn:
            conn.executemany('''INSERT INTO life_ledger_entries(user_id,kind,amount_cents,date,category,note,
                merchant,operation_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)''',
                [(actor, 'expense', 12345, date, changes.get('category', 'Food'), changes.get('note', 'Synthetic note'),
                  changes.get('merchant', 'Test shop'), uuid4().hex, utcnow()) for _ in range(count)])

    def enrich_payload(self):
        token = uuid4().hex
        ledger.save_receipt_draft(self.store, 2, self.candidate(), token)
        identifier = receipt_wines.queue_lookup(self.store, 2, token, 0)
        return self.store.one('SELECT kind FROM jobs WHERE id=?', (identifier,)), json.loads(
            self.store.one('SELECT payload FROM jobs WHERE id=?', (identifier,))['payload'])

    def test_wine_enrichment_is_private_idempotent_and_never_confirms_receipt(self):
        from tests.test_receipt_wines import candidate, lookup_response
        job, payload = self.enrich_payload()
        with patch('thesaurus.wine_enrichment.lookup', return_value=lookup_response(candidate(), cached=True)) as lookup:
            self.run_job(job, payload)
            deliveries = len(self.store.query('SELECT * FROM deliveries'))
            self.run_job(job, payload)
        lookup.assert_called_once()
        quality = json.loads(self.report(payload)['quality'])
        self.assertEqual((quality['feature'], quality['actor_id'], quality['owner_id']), ('life_finance', 2, 2))
        self.assertIn('wine_cellar', quality['required_features'])
        self.assertIn(f'rc:item:{payload["receipt_token"]}:0', str(quality['life_actions']))
        self.assertIn(f'rc:info:{payload["receipt_token"]}:0', str(quality['life_actions']))
        self.assertEqual([row['recipient'] for row in self.store.query('SELECT DISTINCT recipient FROM deliveries')], ['2'])
        self.assertEqual(len(self.store.query('SELECT * FROM deliveries')), deliveries)
        self.assert_no_financial_changes()

    def test_enrichment_result_includes_before_after_names_and_attributes(self):
        from tests.test_receipt_wines import candidate, lookup_response
        job, payload = self.enrich_payload()
        with patch('thesaurus.wine_enrichment.lookup', return_value=lookup_response(candidate(name='Merlot Reserva'))):
            self.run_job(job, payload)
        text = self.report(payload)['text']
        self.assertIn('Название: Wine Test 2020 750 ml → Merlot Reserva', text)
        self.assertIn('Производитель: не указано → Bodega Ejemplo', text)
        self.assertIn('Страна: не указано → Испания', text)
        self.assertNotIn('Год: 2020 →', text)
        self.assert_no_financial_changes()

    def test_wine_enrichment_dry_run_never_calls_http_or_changes_review_state(self):
        from tests.test_receipt_wines import candidate, lookup_response
        job, payload = self.enrich_payload()
        before = self.store.get_setting(f'receipt_ui:2:{payload["receipt_token"]}')
        with patch('thesaurus.wine_enrichment.lookup', return_value=lookup_response(candidate())) as lookup:
            self.run_job(job, payload, dry_run=True)
            lookup.assert_not_called()
            self.assertEqual(self.store.get_setting(f'receipt_ui:2:{payload["receipt_token"]}'), before)
            self.assertFalse(self.store.query('SELECT * FROM reports'))
            self.assertFalse(self.store.query('SELECT * FROM deliveries'))
            self.run_job(job, payload)
        lookup.assert_called_once()
        self.assert_no_financial_changes()

    def test_wine_feature_revocation_cancels_queued_enrichment_before_provider(self):
        _, payload = self.enrich_payload()
        change(self.store, 1, 2, 'wine_cellar', False)
        with patch('thesaurus.wine_enrichment.lookup') as lookup:
            Worker(self.store, self.settings, dry_run=False, charts=False).process_next()
        lookup.assert_not_called()
        state = self.store.one("SELECT state FROM jobs WHERE kind='life_wine_enrich'")['state']
        self.assertEqual(state, 'cancelled')
        self.assertIsNone(self.report(payload))
        self.assert_no_financial_changes()

    def test_wine_feature_revocation_during_enrichment_prevents_report_and_delivery(self):
        from tests.test_receipt_wines import candidate, lookup_response
        job, payload = self.enrich_payload()
        def revoke(*args, **kwargs):
            change(self.store, 1, 2, 'wine_cellar', False)
            return lookup_response(candidate())
        with patch('thesaurus.wine_enrichment.lookup', side_effect=revoke), self.assertRaises(PermissionError):
            self.run_job(job, payload)
        self.assertFalse(self.store.query('SELECT * FROM reports'))
        self.assertFalse(self.store.query('SELECT * FROM deliveries'))
        self.assert_no_financial_changes()

    def test_wine_enrichment_result_is_inaccessible_after_required_feature_revocation(self):
        from tests.test_receipt_wines import candidate, lookup_response
        from thesaurus.telegram_data_jobs import authorized_report
        job, payload = self.enrich_payload()
        with patch('thesaurus.wine_enrichment.lookup', return_value=lookup_response(candidate())):
            self.run_job(job, payload)
        change(self.store, 1, 2, 'wine_cellar', False)
        with self.assertRaises(PermissionError):
            authorized_report(self.store, 2, payload['report_id'])
        client = Mock()
        self.assertEqual(deliver_pending(self.store, client), 0)
        client.send_message.assert_not_called()
        self.assertTrue(all(row['state'] == 'cancelled' for row in self.store.query('SELECT state FROM deliveries')))
        self.assert_no_financial_changes()

    def test_wine_enrichment_payload_cannot_read_another_owner_receipt(self):
        job, payload = self.enrich_payload()
        with patch('thesaurus.wine_enrichment.lookup') as lookup:
            for changes in ({'owner_id': 3}, {'actor_id': 3, 'owner_id': 3}, {'actor_id': 1, 'owner_id': 2}):
                with self.subTest(changes=changes), self.assertRaises((ValueError, PermissionError)):
                    self.run_job(job, {**payload, **changes})
        lookup.assert_not_called()
        self.assertFalse(self.store.query('SELECT * FROM reports'))
        self.assert_no_financial_changes()

    def test_recognition_only_creates_private_review_draft_and_delivery_buttons(self):
        job, payload = self.payload(path=str(self.receipt()))
        with patch('thesaurus.receipts.recognize_receipt', return_value=self.candidate()) as recognize:
            self.run_job(job, payload)
        self.assert_no_financial_changes()
        draft = ledger.get_receipt_draft(self.store, 2, payload['request_id'])
        self.assertEqual(draft['result']['total'], 250)
        self.assertEqual(len(draft['result']['items']), 2)
        self.assertNotIn('PRIVATE FULL OCR', str(draft))
        with self.assertRaises(ValueError):
            ledger.get_receipt_draft(self.store, 3, payload['request_id'])
        report = self.report(payload)
        self.assertEqual(report['section'], 'life')
        self.assertIn('Расходы и запас ещё не изменены', report['text'])
        quality = json.loads(report['quality'])
        self.assertEqual((quality['feature'], quality['actor_id'], quality['owner_id']), ('life_finance', 2, 2))
        self.assertIn('rc:review:' + payload['request_id'], str(quality['life_actions']))
        self.assertEqual([r['recipient'] for r in self.store.query('SELECT DISTINCT recipient FROM deliveries')], ['2'])
        self.assertIn('rc:review:' + payload['request_id'], str(self.store.query('SELECT payload FROM deliveries')))
        recognize.assert_called_once_with(self.receipt(), dry_run=False)

    def test_unavailable_recognition_has_manual_fallback_and_no_draft(self):
        job, payload = self.payload(path=str(self.receipt()))
        with patch('thesaurus.receipts.recognize_receipt', return_value={'status': 'unavailable', 'warnings': ['Нужен локальный OCR']}):
            self.run_job(job, payload)
        self.assert_no_financial_changes()
        self.assertFalse(self.store.query('SELECT * FROM life_receipt_drafts'))
        self.assertIn('lf:add:expense', self.report(payload)['quality'])
        self.assertIn('Нужен локальный OCR', self.report(payload)['text'])

    def test_recognition_rejects_outside_another_owner_and_missing_files_before_ocr(self):
        outside = self.root / 'outside.pdf'
        outside.write_bytes(b'Other private file')
        for path in (outside, self.receipt(3), self.settings.data_dir / 'receipts' / '2' / 'missing.pdf'):
            job, payload = self.payload(path=str(path))
            with patch('thesaurus.receipts.recognize_receipt') as recognize, self.assertRaises(ValueError):
                self.run_job(job, payload)
            recognize.assert_not_called()
        self.assertFalse(self.store.query('SELECT * FROM reports'))

    def test_duplicate_job_reuses_report_without_ocr_or_delivery_duplicates(self):
        job, payload = self.payload(path=str(self.receipt()))
        with patch('thesaurus.receipts.recognize_receipt', return_value=self.candidate()) as recognize:
            self.run_job(job, payload)
            deliveries = len(self.store.query('SELECT * FROM deliveries'))
            self.run_job(job, payload)
        self.assertEqual(recognize.call_count, 1)
        self.assertEqual(len(self.store.query('SELECT * FROM reports')), 1)
        self.assertEqual(len(self.store.query('SELECT * FROM life_receipt_drafts')), 1)
        self.assertEqual(len(self.store.query('SELECT * FROM deliveries')), deliveries)

    def test_report_id_reuse_cannot_cross_actor_owner_or_feature(self):
        job, payload = self.payload('life_export', format='csv')
        self.seed_entries()
        self.run_job(job, payload)
        for changes in ({'actor_id': 3, 'owner_id': 3}, {'actor_id': 1, 'owner_id': 2}, {'owner_id': 3}):
            with self.assertRaises(PermissionError):
                self.run_job(job, {**payload, **changes})
        self.store.execute('UPDATE reports SET quality=? WHERE id=?',
                           (dumps({'feature': 'wine_cellar', 'actor_id': 2, 'owner_id': 2}), payload['report_id']))
        with self.assertRaises(PermissionError):
            self.run_job(job, payload)
        self.assertEqual([r['recipient'] for r in self.store.query('SELECT DISTINCT recipient FROM deliveries')], ['2'])

    def test_disabled_or_unassigned_user_cannot_recognize_or_export(self):
        for actor in (2, 4):
            if actor == 2:
                self.store.execute('UPDATE users SET enabled=0 WHERE user_id=2')
            job, payload = self.payload(path=str(self.receipt(actor)), actor_id=actor, owner_id=actor)
            with patch('thesaurus.receipts.recognize_receipt') as recognize, self.assertRaises(PermissionError):
                self.run_job(job, payload)
            recognize.assert_not_called()
        self.assertFalse(self.store.query('SELECT * FROM reports'))

    def test_revocation_during_ocr_prevents_draft_report_and_delivery(self):
        job, payload = self.payload(path=str(self.receipt()))
        def revoke(*args, **kwargs):
            self.store.set_setting(key(2), [])
            return self.candidate()
        with patch('thesaurus.receipts.recognize_receipt', side_effect=revoke), self.assertRaises(PermissionError):
            self.run_job(job, payload)
        for table in ('life_receipt_drafts', 'reports', 'deliveries'):
            self.assertFalse(self.store.query('SELECT * FROM ' + table))

    def test_worker_cancels_revoked_queued_job_before_recognition(self):
        _, payload = self.payload(path=str(self.receipt()))
        identifier = self.store.enqueue('life_receipt', payload)
        self.store.set_setting(key(2), [])
        with patch('thesaurus.receipts.recognize_receipt') as recognize:
            Worker(self.store, self.settings, dry_run=True, charts=False).process_next()
        recognize.assert_not_called()
        self.assertEqual(self.store.one('SELECT state FROM jobs WHERE id=?', (identifier,))['state'], 'cancelled')

    def test_revocation_after_first_part_blocks_csv_and_followup_buttons(self):
        self.seed_entries()
        self.run_job(*self.payload('life_export', format='csv'))
        client = Mock()
        def revoke(*args, **kwargs):
            change(self.store, 1, 2, 'life_finance', False)
            return SimpleNamespace(id=55)
        client.send_message.side_effect = revoke
        self.assertEqual(deliver_pending(self.store, client), 1)
        client.send_document.assert_not_called()
        self.assertEqual([r['state'] for r in self.store.query('SELECT state FROM deliveries ORDER BY part_index')],
                         ['sent', 'cancelled', 'cancelled'])

    def test_csv_exports_all_rows_of_only_owner_and_selected_dates(self):
        self.seed_entries(1005)
        self.seed_entries(actor=3, note='PRIVATE OTHER OWNER')
        self.seed_entries(date='2025-12-31', note='OUTSIDE PERIOD')
        job, payload = self.payload('life_export', format='csv', start='2026-01-01', end='2026-01-31')
        self.run_job(job, payload)
        with self.document(payload).open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(len(rows), 1006)
        self.assertEqual(rows[1][2:4], ['123.45', 'RUB'])
        self.assertNotIn('PRIVATE OTHER OWNER', str(rows))
        self.assertNotIn('OUTSIDE PERIOD', str(rows))
        self.assertIn('Операций: 1005', self.report(payload)['text'])
        self.assertTrue(self.document(payload).is_relative_to(self.settings.reports_dir.resolve()))

    def test_export_excludes_cancelled_entries(self):
        self.seed_entries(2)
        identifier = ledger.entries(self.store, 2)[0]['id']
        ledger.void_entry(self.store, 2, identifier)
        job, payload = self.payload('life_export', format='csv')
        self.run_job(job, payload)
        with self.document(payload).open(encoding='utf-8-sig', newline='') as stream:
            self.assertEqual(len(list(csv.reader(stream))), 2)

    def test_csv_and_excel_escape_untrusted_formula_fields(self):
        self.seed_entries(category='=1+1', note='\t@HYPERLINK("evil")', merchant=' +SUM(A1:A2)')
        for format_ in ('csv', 'xlsx'):
            job, payload = self.payload('life_export', format=format_)
            self.run_job(job, payload)
            if format_ == 'csv':
                with self.document(payload).open(encoding='utf-8-sig', newline='') as stream:
                    row = list(csv.reader(stream))[1]
            else:
                workbook = load_workbook(self.document(payload), read_only=True, data_only=False)
                try:
                    cells = list(workbook.active.iter_rows())[1]
                    self.assertTrue(all(cell.data_type != 'f' for cell in cells))
                    row = [cell.value for cell in cells]
                finally:
                    workbook.close()
            self.assertEqual(row[4], "'=1+1")
            self.assertEqual(row[5], "'\t@HYPERLINK(\"evil\")")
            self.assertEqual(row[6], "' +SUM(A1:A2)")

    def test_export_uses_consistent_snapshot_when_ledger_changes_mid_write(self):
        self.seed_entries(1005)
        changed = False
        def insert_during_export(value):
            nonlocal changed
            if not changed:
                changed = True
                ledger.add_entry(self.store, 2, 'expense', 999, date='2026-01-11', note='AFTER SNAPSHOT')
            return safe_cell(value)
        job, payload = self.payload('life_export', format='csv')
        with patch.object(jobs, 'safe_cell', side_effect=insert_during_export):
            self.run_job(job, payload)
        with self.document(payload).open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(len(rows), 1006)
        self.assertNotIn('AFTER SNAPSHOT', str(rows))
        self.assertEqual(len(ledger.entries(self.store, 2, limit=1000)), 1000)

    def test_revocation_during_export_removes_artifact_and_prevents_report(self):
        self.seed_entries()
        changed = False
        def revoke(value):
            nonlocal changed
            if not changed:
                changed = True
                self.store.set_setting(key(2), [])
            return safe_cell(value)
        job, payload = self.payload('life_export', format='csv')
        with patch.object(jobs, 'safe_cell', side_effect=revoke), self.assertRaises(PermissionError):
            self.run_job(job, payload)
        self.assertFalse(self.report(payload))
        self.assertFalse(list(self.settings.reports_dir.rglob('*.csv')))
        self.assertFalse(self.store.query('SELECT * FROM deliveries'))

    def test_dry_run_does_not_save_drafts_reports_expenses_or_delivery_and_real_retry_works(self):
        job, payload = self.payload(path=str(self.receipt()))
        with patch('thesaurus.receipts.recognize_receipt', return_value=self.candidate()) as recognize:
            self.run_job(job, payload, dry_run=True)
            self.assertEqual(recognize.call_args.kwargs, {'dry_run': True})
            self.assert_no_financial_changes()
            for table in ('life_receipt_drafts', 'reports', 'deliveries'):
                self.assertFalse(self.store.query('SELECT * FROM ' + table))
            self.run_job(job, payload)
            self.assertEqual(recognize.call_count, 2)
        self.assertIsNotNone(self.report(payload))
        self.assert_no_financial_changes()

    def test_dry_run_export_does_not_create_file_or_delivery(self):
        self.seed_entries()
        self.run_job(*self.payload('life_export', format='xlsx'), dry_run=True)
        self.assertFalse(self.store.query('SELECT * FROM deliveries'))
        self.assertFalse(self.store.query('SELECT * FROM reports'))
        self.assertFalse(list(self.root.rglob('*.xlsx')))

    def test_unknown_kind_bad_identifier_and_invalid_format_cannot_write_artifacts(self):
        for job, payload in (self.payload('life_unknown'), self.payload('life_export', format='html'),
                             self.payload('life_export', format='csv', report_id='../escape')):
            with self.assertRaises(ValueError):
                self.run_job(job, payload)
        self.assertFalse(self.store.query('SELECT * FROM reports'))
        self.assertFalse(self.store.query('SELECT * FROM deliveries'))


if __name__ == '__main__':
    unittest.main()
