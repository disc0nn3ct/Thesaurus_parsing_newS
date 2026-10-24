"""Private sequential wine jobs, resumable prices, and inspectable safe exports."""
import csv
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from uuid import uuid4

from openpyxl import load_workbook

from thesaurus import wine_cellar as cellar
from thesaurus import wine_jobs as jobs
from thesaurus.config import Settings
from thesaurus.delivery import deliver_pending
from thesaurus.feature_access import key
from thesaurus.service import Worker
from thesaurus.store import Store, utcnow


class WineJobTests(unittest.TestCase):
    def test_excel_keeps_all_purchases_and_frozen_consumption_prices_after_refresh(self):
        lot = self.lot()
        cellar.consume_bottles(self.store, 2, lot['id'], 2, price=1250, operation_id='drink-export')
        cellar.record_price(self.store, 2, lot['id'], 2000)
        job, payload = self.payload('wine_cellar_export', format='xlsx')
        self.run_job(job, payload)
        path = Path(json.loads(self.report(payload)['parts'])[0]['path'])
        book = load_workbook(path)
        self.addCleanup(book.close)
        self.assertEqual(book['Моя винотека'].max_row, 2)
        rows = list(book['Выпито и результат'].values)
        self.assertEqual(rows[1][3], 2)
        self.assertEqual(rows[1][6], 1250)
        self.assertEqual(rows[1][7], 500)

    def setUp(self):
        parent = Path(__file__).resolve().parents[1] / '.test-tmp'
        parent.mkdir(exist_ok=True)
        directory = tempfile.TemporaryDirectory(dir=parent)
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.store = Store(self.root/'state.sqlite3')
        self.settings = Settings(root=self.root)
        self.store.add_user(1, 'admin')
        for actor in (2, 3):
            self.store.add_user(actor, subscriptions=['vivino'])
        self.store.add_user(4)
        self.store.set_setting('wine_price_allowed_hosts', ['shop.example'])
        self.enterContext(patch('requests.sessions.Session.request', side_effect=AssertionError('No network in wine jobs tests')))
        self.sequence = 100

    def lot(self, actor=2, **changes):
        return cellar.add_bottle(self.store, actor, {'name': 'Chateau Test', 'winery': 'Domaine Test',
                                 'vintage': 2020, 'volume_ml': 750, 'quantity': 2,
                                 'purchase_price': 1000, **changes})

    def payload(self, kind, **changes):
        self.sequence += 1
        return {'id': self.sequence, 'kind': kind}, {'actor_id': 2, 'report_id': uuid4().hex, **changes}

    def run_job(self, job, payload, **kwargs):
        return jobs.run_job(self.store, self.settings, job, payload, **kwargs)

    def report(self, payload):
        return self.store.one('SELECT * FROM reports WHERE id=?', (payload['report_id'],))

    def quote(self, amount=1500, **changes):
        return {'status': 'ok', 'amount': amount, 'currency': 'RUB', 'as_of': utcnow(),
                'availability': 'InStock', 'title': 'Chateau Test 2020 750 ml',
                'name': 'Chateau Test', 'winery': 'Domaine Test', 'vintage': 2020,
                'volume_ml': 750, 'warnings': [], **changes}

    def link(self, lot, suffix='a'):
        return cellar.add_price_link(self.store, 2, lot['id'], 'https://shop.example/wine/'+suffix)

    def photo(self, actor=2):
        path = self.settings.data_dir/'wine_photos'/str(actor)/'test.jpg'
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b'synthetic photo; mocked OCR')
        return path

    def test_recognition_creates_private_draft_without_adding_bottle(self):
        token = uuid4().hex
        job, payload = self.payload('wine_recognize', path=str(self.photo()), request_id=token)
        draft = {'status': 'candidate', 'candidate': {'name': 'Chateau Test', 'winery': 'Domaine Test',
                 'vintage': 2020}, 'provider': 'mock-local-ocr', 'warnings': ['Проверьте этикетку.']}
        with patch('thesaurus.wine_tools.recognize_label', return_value=draft) as reader:
            self.run_job(job, payload)
        self.assertFalse(cellar.lots(self.store, 2))
        saved = self.store.get_setting('wine_recognition:2:'+token)
        self.assertEqual(saved['candidate']['name'], 'Chateau Test')
        self.assertGreater(saved['expires'], utcnow())
        self.assertIn('бутылка ещё не добавлена', self.report(payload)['text'])
        quality = json.loads(self.report(payload)['quality'])
        self.assertEqual(quality['feature'], 'wine_cellar')
        self.assertIn('wc:recognized:'+token, str(quality['wine_actions']))
        self.assertEqual(reader.call_args.args[0], self.photo())
        self.assertEqual([row['recipient'] for row in self.store.query('SELECT DISTINCT recipient FROM deliveries')], ['2'])

    def test_recognition_unavailable_has_manual_button_and_no_fake_stock(self):
        job, payload = self.payload('wine_recognize', path=str(self.photo()))
        with patch('thesaurus.wine_tools.recognize_label', return_value={'status': 'unavailable', 'candidate': {}, 'warnings': ['Введите вручную.']}):
            self.run_job(job, payload)
        report = self.report(payload)
        self.assertIn('Запас не изменён', report['text'])
        self.assertIn('wc:manual', report['quality'])
        self.assertFalse(cellar.lots(self.store, 2))

    def test_recognition_rejects_other_users_and_outside_paths_before_backend(self):
        for path in (self.photo(3), self.root/'outside.jpg'):
            path.write_bytes(b'other private photograph')
            job, payload = self.payload('wine_recognize', path=str(path))
            with patch('thesaurus.wine_tools.recognize_label') as reader, self.assertRaises(ValueError):
                self.run_job(job, payload)
            reader.assert_not_called()
        self.assertFalse(self.store.query('SELECT * FROM reports'))

    def test_recognition_retry_reuses_report_without_reinvoking_ocr(self):
        job, payload = self.payload('wine_recognize', path=str(self.photo()))
        with patch('thesaurus.wine_tools.recognize_label', return_value={'status': 'unavailable', 'warnings': []}) as reader:
            self.run_job(job, payload)
            initial_parts = len(self.store.query('SELECT * FROM deliveries'))
            self.run_job(job, payload)
        self.assertEqual(reader.call_count, 1)
        self.assertEqual(len(self.store.query('SELECT * FROM reports')), 1)
        self.assertEqual(len(self.store.query('SELECT * FROM deliveries')), initial_parts)

    def test_price_refresh_chooses_lowest_matched_offer_and_keeps_evidence(self):
        lot = self.lot()
        first, second = self.link(lot, 'first'), self.link(lot, 'second')
        job, payload = self.payload('wine_price_refresh')
        offers = {first['url']: self.quote(1700), second['url']: self.quote(1450, availability='LimitedAvailability')}
        with patch('thesaurus.wine_tools.fetch_price', side_effect=lambda url, **kwargs: offers[url]) as fetch:
            self.run_job(job, payload)
        price = cellar.get_lot(self.store, 2, lot['id'])['latest_price']
        self.assertEqual(price['amount'], 1450)
        self.assertEqual(price['url'], second['url'])
        self.assertEqual(price['source'], 'shop_min')
        self.assertEqual(price['metadata']['availability'], 'LimitedAvailability')
        self.assertEqual(price['metadata']['checked_offer_count'], 2)
        self.assertEqual(len(price['metadata']['candidate_offers']), 2)
        self.assertEqual(cellar.get_lot(self.store, 2, lot['id'])['quantity'], 2)
        self.assertEqual(fetch.call_count, 2)
        self.assertEqual(fetch.call_args.kwargs['expected']['vintage'], '2020')
        self.assertEqual(fetch.call_args.kwargs['allowed_hosts'], ['shop.example'])

    def test_quote_after_moscow_midnight_is_dated_as_purchase_day(self):
        with patch.object(cellar, '_today', return_value=date(2026, 1, 2)):
            lot = self.lot(purchase_date='2026-01-02')
            self.link(lot)
            quote = self.quote(1500, as_of='2026-01-01T21:30:00+00:00')
            with patch('thesaurus.wine_tools.fetch_price', return_value=quote):
                self.run_job(*self.payload('wine_price_refresh'))
            latest = cellar.get_lot(self.store, 2, lot['id'])['latest_price']
        self.assertEqual((latest['amount'], latest['price_date']), (1500, '2026-01-02'))
        self.assertEqual(latest['metadata']['observed_at'], quote['as_of'])

    def test_recent_quotes_are_reused_without_network_and_old_cache_is_refreshed(self):
        lot = self.lot()
        link = self.link(lot)
        with patch('thesaurus.wine_tools.fetch_price', return_value=self.quote(1500)) as fetch:
            for _ in range(2):
                self.run_job(*self.payload('wine_price_refresh'))
        self.assertEqual(fetch.call_count, 1)
        cached = self.store.get_setting(f"wine_quote_cache:2:{link['id']}")
        cached['checked_at'] = '2000-01-01T00:00:00+00:00'
        self.store.set_setting(f"wine_quote_cache:2:{link['id']}", cached)
        with patch('thesaurus.wine_tools.fetch_price', return_value=self.quote(1600)) as fetch:
            self.run_job(*self.payload('wine_price_refresh'))
        fetch.assert_called_once()
        self.assertEqual(cellar.get_lot(self.store, 2, lot['id'])['latest_price']['amount'], 1600)

    def test_revoked_shop_does_not_reuse_recent_quote(self):
        lot = self.lot()
        self.link(lot)
        with patch('thesaurus.wine_tools.fetch_price', return_value=self.quote(1500)):
            self.run_job(*self.payload('wine_price_refresh'))
        count = len(self.store.query('SELECT * FROM wine_prices'))
        self.store.set_setting('wine_price_allowed_hosts', [])
        with patch('thesaurus.wine_tools.fetch_price') as fetch:
            job, payload = self.payload('wine_price_refresh')
            self.run_job(job, payload)
        fetch.assert_not_called()
        self.assertEqual(len(self.store.query('SELECT * FROM wine_prices')), count)
        self.assertIn('без обновления: 1', self.report(payload)['text'])

    def test_failed_or_mismatched_quote_preserves_previous_value_and_other_lots(self):
        first = self.lot()
        second = self.lot(name='Chateau Other', winery='Domaine Other')
        link = self.link(first)
        self.link(second, 'other')
        def fetch(url, **kwargs):
            if url == link['url']:
                return {'status': 'unavailable', 'warnings': ['Магазин недоступен.']}
            return self.quote(1800, name='Chateau Other', title='Chateau Other 2020 750 ml', winery='Domaine Other')
        with patch('thesaurus.wine_tools.fetch_price', side_effect=fetch):
            self.run_job(*self.payload('wine_price_refresh'))
        self.assertEqual(cellar.get_lot(self.store, 2, first['id'])['latest_price']['amount'], 1000)
        self.assertEqual(cellar.get_lot(self.store, 2, second['id'])['latest_price']['amount'], 1800)

    def test_dry_run_does_not_fetch_quote_or_change_stock_prices_or_deliver(self):
        lot = self.lot()
        link = self.link(lot)
        before = len(self.store.query('SELECT * FROM wine_prices'))
        job, payload = self.payload('wine_price_refresh')
        with patch('thesaurus.wine_tools._download') as network:
            self.run_job(job, payload, dry_run=True)
        network.assert_not_called()
        self.assertEqual(len(self.store.query('SELECT * FROM wine_prices')), before)
        self.assertIsNone(self.store.get_setting(f"wine_quote_cache:2:{link['id']}"))
        self.assertFalse(self.store.query('SELECT * FROM deliveries'))
        self.assertEqual(cellar.get_lot(self.store, 2, lot['id'])['quantity'], 2)

    def test_partial_price_job_retry_does_not_duplicate_valuation_events(self):
        lot = self.lot()
        self.link(lot)
        job, payload = self.payload('wine_price_refresh')
        with patch('thesaurus.wine_tools.fetch_price', return_value=self.quote(1500)) as fetch:
            with patch.object(jobs, '_result', side_effect=RuntimeError('simulated crash')), self.assertRaises(RuntimeError):
                self.run_job(job, payload)
            self.run_job(job, payload)
        self.assertEqual(fetch.call_count, 1)
        self.assertEqual(len(self.store.query('SELECT * FROM wine_prices')), 2)
        self.assertEqual(len(cellar.valuation_history(self.store, 2)), 2)
        self.assertEqual(len(self.store.query('SELECT DISTINCT report_id,recipient FROM deliveries')), 1)

    def test_access_revoked_during_refresh_stops_later_links_and_value_writes(self):
        lot = self.lot()
        self.link(lot, 'one')
        self.link(lot, 'two')
        def revoke(url, **kwargs):
            self.store.set_setting(key(2), [])
            return self.quote()
        with patch('thesaurus.wine_tools.fetch_price', side_effect=revoke) as fetch, self.assertRaises(PermissionError):
            self.run_job(*self.payload('wine_price_refresh'))
        self.assertEqual(fetch.call_count, 1)
        self.assertEqual(len(self.store.query('SELECT * FROM wine_prices')), 1)
        self.assertFalse(self.store.query('SELECT * FROM reports'))

    def test_csv_and_excel_have_russian_price_headers_and_formula_safe_text(self):
        self.lot(name='=1+1', winery='@SUM(A1:A9)', notes='-formula')
        for format_ in ('csv', 'xlsx'):
            job, payload = self.payload('wine_cellar_export', format=format_)
            self.run_job(job, payload)
            path = Path(json.loads(self.report(payload)['parts'])[0]['path'])
            self.assertTrue(path.is_relative_to(self.settings.reports_dir))
            if format_ == 'csv':
                with path.open(encoding='utf-8-sig', newline='') as stream:
                    rows = list(csv.reader(stream))
                self.assertEqual(rows[1][0], "'=1+1")
                self.assertTrue(any('Цена покупки' in value for value in rows[0]))
            else:
                book = load_workbook(path)
                self.addCleanup(book.close)
                rows = list(book['Моя винотека'].values)
                self.assertEqual(rows[1][0], "'=1+1")
                self.assertTrue(any('Цена покупки' in str(value) for value in rows[0]))
                self.assertIn('История оценки', book.sheetnames)
                self.assertNotEqual(book['Моя винотека']['A2'].data_type, 'f')

    def test_html_export_escapes_untrusted_text_and_has_no_external_resources(self):
        self.lot(name='<img src=x onerror=alert(1)>', notes='<script>PRIVATE</script>')
        job, payload = self.payload('wine_cellar_export', format='html')
        self.run_job(job, payload)
        path = Path(json.loads(self.report(payload)['parts'])[0]['path'])
        page = path.read_text(encoding='utf-8')
        self.assertIn('&lt;img', page)
        self.assertIn('&lt;script&gt;PRIVATE&lt;/script&gt;', page)
        self.assertNotIn('<script>PRIVATE', page)
        self.assertNotIn('<script src=', page)
        self.assertIn('Найти вино', page)

    def test_chart_renders_unknown_values_without_filling_them_with_zero(self):
        self.lot(purchase_price=None)
        job, payload = self.payload('wine_cellar_export', format='chart')
        self.run_job(job, payload)
        part = json.loads(self.report(payload)['parts'])[0]
        self.assertEqual(part['kind'], 'photo')
        self.assertGreater(Path(part['path']).stat().st_size, 1000)
        self.assertIsNone(cellar.portfolio(self.store, 2)['current_value'])

    def test_empty_chart_job_delivers_guidance_without_a_blank_image(self):
        job, payload = self.payload('wine_cellar_export', format='chart')
        self.run_job(job, payload)
        report = self.report(payload)
        self.assertEqual(json.loads(report['parts']), [])
        self.assertIn('нет сохранённых добавлений бутылок', report['text'])
        self.assertIn('wc:menu', report['quality'])
        self.assertIn('wc:add', report['quality'])
        self.assertFalse(list(self.settings.reports_dir.rglob('*.png')))

    def test_chart_uses_final_daily_stock_in_moscow_instead_of_receipt_milliseconds(self):
        snapshots = [
            {'created_at': '2026-10-10T20:59:59+00:00', 'quantity': 1, 'current_value': 600, 'purchase_cost': 600},
            {'created_at': '2026-10-10T21:00:00+00:00', 'quantity': 2, 'current_value': 1200, 'purchase_cost': 1200},
            {'created_at': '2026-10-10T21:00:00.005+00:00', 'quantity': 7, 'current_value': 4433.44, 'purchase_cost': 4433.44},
        ]
        figures = []
        with patch.object(cellar, 'valuation_history', return_value=snapshots), \
                patch('matplotlib.figure.Figure.savefig', autospec=True, side_effect=lambda figure, *args, **kwargs: figures.append(figure)):
            jobs._chart(self.store, self.settings, 2, 2, uuid4().hex)
        values, _results, stock = figures[0].axes
        self.assertEqual(list(values.lines[0].get_ydata()), [600, 4433.44])
        self.assertEqual(list(stock.lines[0].get_ydata()), [1, 7])
        self.assertEqual([day.strftime('%d.%m.%Y') for day in values.lines[0].get_xdata()], ['10.10.2026', '11.10.2026'])
        self.assertIn('Москва', stock.get_xlabel())

    def test_fully_consumed_stock_still_has_a_chart_of_preserved_history(self):
        lot = self.lot()
        cellar.adjust_stock(self.store, 2, lot['id'], -2, reason='drink', operation_id='consumed-for-chart')
        self.assertFalse(cellar.lots(self.store, 2))
        job, payload = self.payload('wine_cellar_export', format='chart')
        self.run_job(job, payload)
        part = json.loads(self.report(payload)['parts'])[0]
        self.assertEqual(part['kind'], 'photo')
        self.assertTrue(Path(part['path']).is_file())

    def test_only_admin_can_export_other_owner_and_report_is_private_to_requester(self):
        self.lot(name='PRIVATE-OWNER-TWO')
        job, payload = self.payload('wine_cellar_export', actor_id=3, owner_id=2, format='csv')
        with self.assertRaises(PermissionError):
            self.run_job(job, payload)
        job, payload = self.payload('wine_cellar_export', actor_id=1, owner_id=2, format='csv')
        self.run_job(job, payload)
        quality = json.loads(self.report(payload)['quality'])
        self.assertEqual((quality['actor_id'], quality['owner_id']), (1, 2))
        self.assertEqual([row['recipient'] for row in self.store.query('SELECT DISTINCT recipient FROM deliveries')], ['1'])

    def test_reusing_another_actors_report_identifier_cannot_deliver_it(self):
        self.lot(name='PRIVATE-OWNER-TWO')
        job, payload = self.payload('wine_cellar_export', format='csv')
        self.run_job(job, payload)
        second = {**payload, 'actor_id': 3}
        with self.assertRaises(PermissionError):
            self.run_job(job, second)
        self.assertEqual([row['recipient'] for row in self.store.query('SELECT DISTINCT recipient FROM deliveries')], ['2'])

    def test_worker_cancels_revoked_function_before_recognition(self):
        _, payload = self.payload('wine_recognize', path=str(self.photo()))
        job_id = self.store.enqueue('wine_recognize', payload)
        self.store.set_setting(key(2), [])
        with patch('thesaurus.wine_tools.recognize_label') as reader:
            Worker(self.store, self.settings, dry_run=True, charts=False).process_next()
        reader.assert_not_called()
        self.assertEqual(self.store.one('SELECT state FROM jobs WHERE id=?', (job_id,))['state'], 'cancelled')
        self.assertFalse(self.store.query('SELECT * FROM reports'))

    def test_revocation_during_delivery_blocks_private_table_and_later_parts(self):
        from thesaurus.feature_access import change
        self.lot()
        self.run_job(*self.payload('wine_cellar_export', format='csv'))
        client = Mock()
        def revoke(*args, **kwargs):
            change(self.store, 1, 2, 'wine_cellar', False)
            return SimpleNamespace(id=55)
        client.send_message.side_effect = revoke
        self.assertEqual(deliver_pending(self.store, client), 1)
        client.send_document.assert_not_called()
        states = [row['state'] for row in self.store.query('SELECT state FROM deliveries ORDER BY part_index')]
        self.assertEqual(states, ['sent', 'cancelled'])

    def test_bad_report_id_and_unknown_job_kind_cannot_write_artifacts(self):
        job, payload = self.payload('wine_cellar_export', report_id='../escape', format='csv')
        with self.assertRaises(ValueError):
            self.run_job(job, payload)
        with self.assertRaises(ValueError):
            self.run_job(*self.payload('wine_unknown'))
        self.assertFalse(self.store.query('SELECT * FROM reports'))
