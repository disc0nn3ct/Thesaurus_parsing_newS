import csv
import shutil
import unittest
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from thesaurus.exporting import COLUMNS, export_data, safe_cell
from thesaurus.store import Store


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.folder = Path(__file__).resolve().parent / f"export_test_{uuid4().hex}"
        self.folder.mkdir()
        self.store = Store(self.folder / "state.sqlite3")
        self.store.save_observations([
            {"source": "wordstat", "metric": "пожар", "region_id": 225, "frequency": "daily", "date": "2025-01-01", "value": 0, "unit": "count"},
            {"source": "wordstat", "metric": "пожар", "region_id": 1, "frequency": "daily", "date": "2025-01-02", "value": None, "unit": "count"},
            {"source": "wordstat", "metric": "=1+1", "region_id": 225, "frequency": "weekly", "date": "2025-01-06", "value": 2, "unit": "count"},
            {"source": "ruonia", "metric": "rate", "frequency": "daily", "date": "2025-01-01", "value": 12.5, "unit": "percent"},
        ])

    def tearDown(self):
        self.assertEqual(self.folder.resolve().parent, Path(__file__).resolve().parent)
        shutil.rmtree(self.folder)

    def test_csv_filters_preserve_source_unit_quality_and_zero(self):
        path = export_data(self.store, self.folder / "data.csv", date_from="2025-01-01", date_to="2025-01-01", phrases=["пожар"], regions=[225], source="wordstat", frequency="daily")
        with path.open(encoding="utf-8-sig", newline="") as handle:
            data = list(csv.DictReader(handle))
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]["value"], "0.0")
        self.assertEqual(data[0]["unit"], "count")
        self.assertEqual(data[0]["quality"], "ok")
        self.assertEqual(data[0]["source"], "wordstat")

    def test_csv_missing_value_stays_empty(self):
        path = export_data(self.store, self.folder / "missing.csv", regions=[1])
        with path.open(encoding="utf-8-sig", newline="") as handle:
            data = list(csv.DictReader(handle))
        self.assertEqual(data[0]["value"], "")
        self.assertEqual(data[0]["quality"], "missing")

    def test_excel_escapes_formulas_and_retains_numeric_values(self):
        path = export_data(self.store, self.folder / "data.xlsx", format="xlsx", frequency="weekly")
        workbook = load_workbook(path)
        sheet = workbook.active
        self.assertEqual(sheet.cell(2, COLUMNS.index("metric") + 1).value, "'=1+1")
        self.assertEqual(sheet.cell(2, COLUMNS.index("metric") + 1).data_type, "s")
        self.assertEqual(sheet.cell(2, COLUMNS.index("value") + 1).value, 2)
        workbook.close()

    def test_empty_filters_export_headers(self):
        path = export_data(self.store, self.folder / "empty.csv", phrases=[])
        with path.open(encoding="utf-8-sig", newline="") as handle:
            data = list(csv.reader(handle))
        self.assertEqual(data, [COLUMNS])

    def test_sql_like_phrase_remains_parameter(self):
        path = export_data(self.store, self.folder / "injection.csv", phrases=["' OR 1=1 --"])
        with path.open(encoding="utf-8-sig", newline="") as handle:
            self.assertEqual(len(list(csv.reader(handle))), 1)
        self.assertEqual(len(self.store.query("SELECT * FROM observations")), 4)

    def test_invalid_filters_fail_before_file_creation(self):
        for filters in ({"format": "pdf"}, {"date_from": "bad"}, {"date_from": "2025-02-01", "date_to": "2025-01-01"}, {"source": "unknown"}, {"regions": [-1]}, {"frequency": "hourly"}):
            with self.subTest(filters=filters), self.assertRaises(ValueError):
                export_data(self.store, self.folder / "invalid.csv", **filters)
        self.assertFalse((self.folder / "invalid.csv").exists())

    def test_formula_prefix_controls_are_escaped(self):
        for value in ("=1+1", "+SUM(1,2)", "-2+3", "@SUM(1,2)", " \t=1+1", "\ufeff=1+1", "\tplain"):
            with self.subTest(value=value):
                self.assertTrue(safe_cell(value).startswith("'"))
        self.assertEqual(safe_cell(-2.5), -2.5)
        self.assertIsNone(safe_cell(None))


if __name__ == "__main__":
    unittest.main()
