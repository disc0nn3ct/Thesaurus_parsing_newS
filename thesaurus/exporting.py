"""Filtered observation exports; formulas in untrusted string cells are escaped."""

from __future__ import annotations

import csv
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

COLUMNS = ["source", "metric", "region_id", "device", "frequency", "date", "value", "unit",
           "quality", "retrieved_at", "source_url", "keyword_version", "cohort_version"]


def safe_cell(value):
    """Neutralize spreadsheet formula prefixes, including whitespace/control tricks."""
    if not isinstance(value, str):
        return value
    if any(unicodedata.category(char) == "Cc" and char not in "\t\r\n" for char in value):
        value = "".join(char for char in value if unicodedata.category(char) != "Cc" or char in "\t\r\n")
    trimmed = value
    while trimmed and (trimmed[0].isspace() or unicodedata.category(trimmed[0]) == "Cf"):
        trimmed = trimmed[1:]
    if (trimmed and trimmed[0] in "=+-@") or value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def _date(value, label):
    if value is None:
        return None
    try:
        return date.fromisoformat(str(value)).isoformat()
    except ValueError:
        raise ValueError(f"{label} must be an ISO date") from None


def export_data(store, path, format="csv", date_from=None, date_to=None, phrases=None,
                regions=None, source=None, frequency=None, *, forecast_rows=None) -> Path:
    """Export Store observations; an empty selection still produces column headers."""
    format = str(format).lower()
    if format not in {"csv", "xlsx"}:
        raise ValueError("Supported export formats: csv, xlsx")
    start, end = _date(date_from, "date_from"), _date(date_to, "date_to")
    if start and end and start > end:
        raise ValueError("date_from must not be later than date_to")
    conditions, params = [], []
    if start:
        conditions.append("date>=?")
        params.append(start)
    if end:
        conditions.append("date<=?")
        params.append(end)
    if source is not None:
        if source not in {"ruonia", "wordstat"}:
            raise ValueError("Unknown observation source")
        conditions.append("source=?")
        params.append(source)
    if frequency is not None:
        if frequency not in {"daily", "weekly", "monthly"}:
            raise ValueError("Unknown observation frequency")
        conditions.append("frequency=?")
        params.append(frequency)
    for column, values in (("metric", phrases), ("region_id", regions)):
        if values is None:
            continue
        values = [values] if isinstance(values, (str, int)) else list(values)
        if column == "region_id":
            try:
                values = [int(value) for value in values]
            except (TypeError, ValueError):
                raise ValueError("Regions must contain integer IDs") from None
            if any(value <= 0 for value in values):
                raise ValueError("Regions must be positive")
        else:
            values = [str(value) for value in values]
        if len(values) > 500:
            raise ValueError("Too many export filters")
        if not values:
            conditions.append("0")
        else:
            conditions.append(f"{column} IN ({','.join('?' for _ in values)})")
            params.extend(values)
    where = " WHERE " + " AND ".join(conditions) if conditions else ""
    rows = store.query("SELECT " + ",".join(COLUMNS) + " FROM observations" + where + " ORDER BY date,source,metric,region_id,device,frequency", params)
    columns = COLUMNS
    if forecast_rows is not None:
        columns = COLUMNS + ['record_type', 'model', 'lower', 'upper', 'baseline', 'forecast_created_at']
        rows = [{**row, 'record_type': 'actual'} for row in rows] + forecast_rows
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    if format == "csv":
        with output.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(columns)
            writer.writerows([safe_cell(row.get(column)) for column in columns] for row in rows)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Observations"
        sheet.append(columns)
        for row in rows:
            sheet.append([safe_cell(row.get(column)) for column in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244563")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
        workbook.save(output)
    return output
