"""Sequential local wine tools and private, resumable result delivery."""
import csv
import html
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlsplit
from uuid import uuid4
from zoneinfo import ZoneInfo

from . import wine_cellar as cellar
from .delivery import prepare_delivery
from .exporting import safe_cell
from .feature_access import require
from .store import dumps, utcnow
from .telegram_data_jobs import _artifact_path, report_identifier


def _result(store, actor, report_id, text, *, parts=None, buttons=None, owner=None):
    require(store, actor, 'wine_cellar')
    quality = {'feature': 'wine_cellar', 'actor_id': actor, 'owner_id': owner or actor}
    if buttons:
        quality['wine_actions'] = buttons
    store.execute("""INSERT OR IGNORE INTO reports(id,section,created_at,keyword_version,cohort_version,
        text,full_text,quality,parts) VALUES(?,'vivino',?,1,0,?,?,?,?)""",
                  (report_id, utcnow(), text, text, dumps(quality), dumps(parts or [])))


def _recognize(store, settings, job, payload, *, dry_run=False):
    from .wine_tools import recognize_label
    actor = payload['actor_id']
    root = (settings.data_dir/'wine_photos'/str(actor)).resolve()
    path = Path(payload.get('path', '')).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise ValueError('Фото не найдено в вашей винотеке. Отправьте его заново.')
    result = recognize_label(path, dry_run=dry_run)
    require(store, actor, 'wine_cellar')
    token = payload.get('request_id') or uuid4().hex
    report_identifier(token)
    draft = {**result, 'expires': (datetime.now(timezone.utc)+timedelta(days=1)).isoformat()}
    store.set_setting(f'wine_recognition:{actor}:{token}', draft)
    warnings = '\n'.join(str(value)[:500] for value in result.get('warnings', [])[:5])
    if result.get('status') == 'candidate' and result.get('candidate', {}).get('name'):
        candidate = result['candidate']
        text = ('Моя винотека · этикетка\nРаспознавание подготовило черновик. Проверьте название, производителя и год; '
                'бутылка ещё не добавлена.\n\n'+str(candidate['name'])[:300]+'\n'+str(candidate.get('winery') or '')[:200]
                +'\nГод: '+str(candidate.get('vintage') or 'нужно уточнить')+'\n'+warnings)
        buttons = [[{'text': 'Проверить и добавить вино', 'callback_data': f'wc:recognized:{token}'}]]
    else:
        text = ('Моя винотека · этикетка\nАвтоматическое распознавание не выполнено. '
                'Запас не изменён. Можно добавить вино по названию.\n'+warnings)
        buttons = [[{'text': 'Добавить по названию', 'callback_data': 'wc:manual'}]]
    buttons.append([{'text': 'Моя винотека', 'callback_data': 'wc:menu'}])
    _result(store, actor, payload['report_id'], text, buttons=buttons)


def _prices(store, job, payload, *, dry_run=False):
    from .wine_retail import allowed_hosts as retail_hosts
    from .wine_retail import availability_label, profile_for_url, quote_eligible
    from .wine_tools import fetch_price, validate_product_url
    actor = payload['actor_id']
    lots = {row['id']: row for row in cellar.lots(store, actor)}
    links = cellar.price_links(store, actor, payload.get('lot_id'))
    lines, refreshed, cached, errors = [], 0, 0, 0
    matched = {}
    allowed_hosts = retail_hosts(store)
    for link in links[:100]:
        require(store, actor, 'wine_cellar')
        lot = lots.get(link['lot_id'])
        if not lot or lot['quantity'] <= 0:
            continue
        try:
            validate_product_url(link['url'], allowed_hosts=retail_hosts(store))
        except ValueError:
            errors += 1
            lines.append(f"{lot['name']}: магазин больше не разрешён администратором или ссылка недопустима.")
            continue
        profile = profile_for_url(link['url'])
        if profile and not profile['automatic']:
            errors += 1
            lines.append(f"{lot['name']}: {profile['title']} — {profile['reason']} Доступна ручная цена.")
            continue
        cache_key = f"wine_quote_cache:{actor}:{link['id']}"
        previous = store.get_setting(cache_key)
        recent = previous and previous.get('checked_at', '') > (datetime.now(timezone.utc)-timedelta(minutes=30)).isoformat()
        if recent:
            quote = previous['quote']
            cached += 1
        else:
            try:
                quote = fetch_price(link['url'], dry_run=dry_run, expected=lot, allowed_hosts=allowed_hosts)
            except (ValueError, OSError, RuntimeError):
                quote = {'status': 'unavailable', 'warnings': ['Магазин не вернул проверяемую цену.']}
            if not dry_run:
                store.set_setting(cache_key, {'checked_at': utcnow(), 'quote': quote})
        if profile and not dry_run:
            store.set_setting('wine_retail_status:'+profile['id'], {
                'shop_id': profile['id'], 'title': profile['title'], 'city': quote.get('city'),
                'checked_at': quote.get('as_of') or (previous or {}).get('checked_at') or utcnow(),
                'status': quote.get('status', 'unavailable'), 'cached': bool(recent),
                'reason': ('Цена и наличие карточки проверены.' if quote.get('status') == 'ok' else
                           '; '.join(str(item)[:180] for item in quote.get('warnings', [])[:2]))})
        if quote.get('status') != 'ok':
            errors += 1
            explanation = '; '.join(str(item)[:160] for item in quote.get('warnings', [])[:2])
            lines.append(f"{lot['name']}: цена не обновлена ({quote.get('status', 'unavailable')}). {explanation}")
            continue
        final_url = quote.get('url') or quote.get('source') or link['url']
        # Administrator changes and expired offers invalidate a previously cached quote.
        try:
            validate_product_url(link['url'], allowed_hosts=retail_hosts(store))
            final_url = validate_product_url(final_url, allowed_hosts=retail_hosts(store))
        except ValueError:
            errors += 1
            lines.append(f"{lot['name']}: магазин отключён во время проверки; цена не изменена.")
            continue
        final_profile = profile_for_url(final_url)
        if profile and ((urlsplit(final_url).hostname or '').removeprefix('www.') != profile['host']):
            errors += 1
            lines.append(f"{lot['name']}: карточка перенаправлена в другой магазин; цена не изменена.")
            continue
        matched_offer, reasons = quote_eligible(quote, lot, moscow=final_profile is not None)
        if (not matched_offer or quote.get('currency') != 'RUB'
                or quote.get('availability') not in {'InStock', 'LimitedAvailability'}):
            errors += 1
            lines.append(f"{lot['name']}: предложение не подходит вашей бутылке или товар недоступен. " + '; '.join(reasons))
            continue
        matched.setdefault(lot['id'], []).append({**quote, 'url': final_url})
        refreshed += 1
        place = f"; {quote['city']}, {availability_label(quote)}" if quote.get('city') else '; регион магазина не подтверждён'
        lines.append(f"{lot['name']}: {quote['amount']:,.2f} ₽ за бутылку{place}; проверено {str(quote.get('as_of') or utcnow())[:19]}.\n{final_url}")
    if not dry_run:
        for lot_id, offers in matched.items():
            current_hosts = retail_hosts(store)
            eligible = []
            for offer in offers:
                try:
                    validate_product_url(offer['url'], allowed_hosts=current_hosts)
                except ValueError:
                    continue
                eligible.append(offer)
            offers = eligible
            if not offers:
                lines.append(f"{lots[lot_id]['name']}: источники отключены; сохранённая оценка не изменена.")
                continue
            # Unscoped custom links cannot undercut a confirmed Moscow quote.
            scoped = [offer for offer in offers if offer.get('city') == 'Москва' and offer.get('price_kind') == 'public']
            if scoped:
                offers = scoped
            best = min(offers, key=lambda offer: (offer['amount'], offer['url']))
            try:
                observed = datetime.fromisoformat(str(best.get('as_of') or utcnow()))
                if observed.tzinfo is None:
                    observed = observed.replace(tzinfo=timezone.utc)
                cellar.record_price(store, actor, lot_id, best['amount'],
                                    price_date=observed.astimezone(ZoneInfo('Europe/Moscow')).date().isoformat(), source='shop_min', url=best['url'],
                                    operation_id=f"wine-price:{job['id']}:{lot_id}",
                                    title=best.get('title'), vintage=best.get('vintage'), volume_ml=best.get('volume_ml'),
                                    observed_at=best.get('as_of'), currency='RUB', available=True,
                                    availability=best['availability'], price_status=best['status'],
                                    metadata={'selection': 'lowest_matched_available_offer',
                                              'checked_offer_count': len(offers), 'candidate_offers': offers,
                                              **{key: best[key] for key in ('shop_id', 'city', 'price_kind', 'availability_scope',
                                                                          'available_shops', 'shop_address') if key in best}})
            except ValueError:
                errors += 1
                lines.append(f"{lots[lot_id]['name']}: оценка не сохранена — параметры предложения изменились.")
    if not links:
        lines.append('Сначала сохраните ссылку на конкретное вино в карточке бутылки. Магазины разрешает администратор.')
    if len(links) > 100:
        lines.append(f'В этом задании проверены первые 100 ссылок из {len(links)}. Для остальных выберите бутылку отдельно.')
    summary = (f'Моя винотека · цены\nОбновлено предложений: {refreshed}; из кэша последних 30 минут: {cached}; '
               f'без обновления: {errors}.\nДля оценки выбрана минимальная доступная цена среди проверенных предложений '
               'с совпавшими названием, годом и объёмом. При наличии проверенной московской цены '
               'предложения без подтверждённого региона с ней не смешиваются. Доставка и скидки по картам не учитываются. '
               'Это стоимость покупки в магазине, не гарантированная цена продажи вашей бутылки.\n\n'
               +'\n\n'.join(lines))
    _result(store, actor, payload['report_id'], summary, buttons=[[
        {'text': 'Стоимость моей винотеки', 'callback_data': f'wc:value:{actor}'},
        {'text': 'Моя винотека', 'callback_data': 'wc:menu'}]])


def _search(store, payload, *, dry_run=False):
    from .wine_retail import search
    actor, token = payload['actor_id'], payload['report_id']
    result = search(store, actor, payload['lot_id'], dry_run=dry_run)
    store.set_setting(f'wine_retail_search:{actor}:{token}', result)
    lines = ['Цены на вино · Москва', 'Поиск не меняет запас и стоимость. Проверьте найденную бутылку перед подтверждением.']
    if not result['sources']:
        lines.append('Магазины ещё не подключены. Администратор выбирает их в разделе «Вина → Магазины · Москва».')
    for item in result['sources']:
        lines.append(f"{item['title']}: {item['reason']}" + (' Из кэша последних 30 минут.' if item.get('cached') else ''))
    lines.append(f"Найдено предложений с проверенной ценой: {len(result['candidates'])}. "
                 'Название, год и объём должны совпадать; цены по картам и варианты «от» не используются.')
    buttons = [[{'text': 'Посмотреть найденные бутылки', 'callback_data': f'wc:offers:{token}'}]] if result['candidates'] else []
    buttons.append([{'text': 'Цены в Москве', 'callback_data': f"wc:retail:{payload['lot_id']}"}])
    _result(store, actor, token, '\n\n'.join(lines), buttons=buttons)


LABELS = {'id': 'Партия', 'name': 'Вино', 'winery': 'Производитель', 'vintage': 'Год урожая',
          'country': 'Страна', 'region': 'Регион', 'grapes': 'Сорта винограда', 'metadata_sources': 'Источники характеристик',
          'quantity': 'Бутылок', 'volume_ml': 'Объём бутылки, мл', 'purchase_price': 'Цена покупки за бутылку, ₽',
          'purchase_date': 'Дата покупки', 'location': 'Место хранения', 'notes': 'Заметки',
          'current_price': 'Оценка за бутылку, ₽', 'price_date': 'Дата оценки', 'price_source': 'Источник оценки',
          'price_url': 'Источник, ссылка', 'currency': 'Валюта', 'purchase_cost': 'Затраты на запас, ₽',
          'current_value': 'Оценка запаса, ₽', 'created_at': 'Добавлено', 'updated_at': 'Обновлено'}
LABELS.update(purchase_price_rub='Цена покупки за бутылку, ₽', latest_price_rub='Оценка за бутылку, ₽',
              value_rub='Оценка запаса, ₽', purchased_quantity='Всего куплено, шт.', consumed_quantity='Выпито, шт.',
              closed_result='Зафиксированный условный результат, ₽', closed_unknown_quantity='Выпито без расчёта результата, шт.')


def _chart(store, settings, actor, owner, report_id):
    history = cellar.valuation_history(store, actor, owner=owner)
    if not history:
        return None
    import numpy as np
    import pandas as pd
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.dates import (
        AutoDateLocator,
        ConciseDateFormatter,
        DateFormatter,
        DayLocator,
    )
    from matplotlib.figure import Figure
    from matplotlib.ticker import MaxNLocator
    path = _artifact_path(settings, report_id, 'wine-value.png')
    figure = Figure(figsize=(11, 10), layout='constrained')
    FigureCanvasAgg(figure)
    axes, result_axes, stock_axes = figure.subplots(3, 1, sharex=True, gridspec_kw={'height_ratios': [3, 2, 1]})
    if history:
        # A receipt produces several snapshots in one transaction. Plot its final
        # daily state instead of an apparent price rise over a few milliseconds.
        days = pd.to_datetime([row['created_at'] for row in history], utc=True, format='ISO8601').tz_convert('Europe/Moscow').normalize()
        daily = dict(zip(days, history, strict=True))
        dates = pd.DatetimeIndex(sorted(daily))
        history = [daily[day] for day in dates]
        values = [row.get('current_value', row.get('value')) for row in history]
        costs = [row.get('purchase_cost', row.get('cost')) for row in history]
        axes.plot(dates, [np.nan if value is None else value for value in values], 'o-', markersize=3, label='Оценка всего запаса · при полном охвате')
        axes.plot(dates, [np.nan if value is None else value for value in costs], '--s', markersize=6, markerfacecolor='none', label='Затраты на оставшийся запас')
        partial = [row.get('known_current_value') if row.get('current_value', row.get('value')) is None and row.get('priced_quantity', 0) else np.nan for row in history]
        if any(value is not None and np.isfinite(value) for value in partial):
            axes.scatter(dates, partial, marker='x', color='#bd861b', label='Оценена только часть запаса')
        stock_axes.step(dates, [row['quantity'] for row in history], where='post', color='#7a3a53')
        stock_axes.scatter(dates, [row['quantity'] for row in history], color='#7a3a53', s=10)
        stock_axes.set_ylim(0, max(1, max(row['quantity'] for row in history))*1.15)
        axes.legend(loc='best', fontsize=8)
        results = cellar.daily_result_history(store,actor,owner=owner)
        by_day = {row['day']:row for row in results}
        for field,label,style in [('open_result','Открытый · оставшийся запас','o-'),
                                  ('closed_result','Зафиксированный · выпито','s-'),
                                  ('combined_result','Общий результат','--')]:
            values_for_result = [by_day.get(day.date().isoformat(),{}).get(field) for day in dates]
            result_axes.plot(dates,[np.nan if value is None else value for value in values_for_result],style,
                             markersize=3,label=label,drawstyle='steps-post')
        result_axes.axhline(0,color='#888888',linewidth=.8)
        result_axes.legend(loc='best',fontsize=8)
        if not any(value is not None for value in values+costs) and not any(np.isfinite(value) for value in partial if value is not None):
            axes.set_yticks([])
            axes.text(.5, .5, 'Цены пока не указаны. Ниже показан размер запаса.', transform=axes.transAxes, ha='center')
        if len(dates) == 1:
            stock_axes.set_xlim(dates[0]-pd.Timedelta(days=1), dates[0]+pd.Timedelta(days=1))
            stock_axes.xaxis.set_major_locator(DayLocator(tz='Europe/Moscow'))
            stock_axes.xaxis.set_major_formatter(DateFormatter('%d.%m.%Y', tz='Europe/Moscow'))
        else:
            locator = AutoDateLocator(minticks=3, maxticks=7, tz='Europe/Moscow')
            stock_axes.xaxis.set_major_locator(locator)
            stock_axes.xaxis.set_major_formatter(ConciseDateFormatter(locator, tz='Europe/Moscow'))
        figure.autofmt_xdate()
    axes.set_title('Моя винотека · стоимость, условный результат и запас')
    axes.set_ylabel('Рубли')
    result_axes.set_ylabel('Условный результат, ₽')
    stock_axes.set_ylabel('Бутылок')
    stock_axes.set_xlabel('Дата · последнее сохранённое состояние за день (Москва)')
    stock_axes.yaxis.set_major_locator(MaxNLocator(integer=True))
    for ax in (axes, result_axes, stock_axes):
        ax.grid(alpha=.18)
    figure.savefig(path, dpi=145)
    return path


def _export(store, settings, payload):
    actor = payload['actor_id']
    owner = cellar.require_owner(store, actor, payload.get('owner_id'))
    format_ = payload.get('format')
    if format_ not in {'csv', 'xlsx', 'html', 'chart'}:
        raise ValueError('Выберите CSV, Excel, HTML или график.')
    report_id = payload['report_id']
    rows = cellar.export_rows(store, actor, owner=owner)
    if format_ == 'chart':
        path = _chart(store, settings, actor, owner, report_id)
        if path is None:
            buttons = [[{'text': 'Моя винотека', 'callback_data': 'wc:menu'}]]
            if actor == owner:
                buttons.insert(0, [{'text': 'Добавить вино', 'callback_data': 'wc:add'}])
            _result(store, actor, report_id,
                    'График винотеки пока недоступен: нет сохранённых добавлений бутылок. '
                    'Если чек уже загружен, откройте «Моя винотека» и продолжите добавление из чека. '
                    'График появится после подтверждения запасов.', buttons=buttons, owner=owner)
            return
        part = {'kind': 'photo', 'path': str(path), 'caption': 'Стоимость, условный результат и запас · по сохранённым состояниям за день. Поздние цены не меняют прошлый результат; без сопоставимых цен линия прерывается.'}
    else:
        columns = list(dict.fromkeys(key for row in rows for key in row)) or list(LABELS)[:10]
        path = _artifact_path(settings, report_id, 'wine-cellar.'+format_)
        if format_ == 'csv':
            with path.open('w', encoding='utf-8-sig', newline='') as stream:
                writer = csv.writer(stream)
                writer.writerow([safe_cell(LABELS.get(column, column)) for column in columns])
                writer.writerows([[safe_cell(row.get(column)) for column in columns] for row in rows])
        elif format_ == 'xlsx':
            from openpyxl import Workbook
            book = Workbook()
            sheet = book.active
            sheet.title = 'Моя винотека'
            sheet.append([safe_cell(LABELS.get(column, column)) for column in columns])
            for row in rows:
                sheet.append([safe_cell(row.get(column)) for column in columns])
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(55, max(16, max(len(str(cell.value or '')) for cell in column)+2))
            history_sheet = book.create_sheet('История оценки')
            history = cellar.valuation_history(store, actor, owner=owner)
            history_columns = ['created_at', 'quantity', 'purchase_cost', 'current_value', 'priced_quantity', 'known_current_value']
            history_sheet.append(['Дата', 'Бутылок', 'Затраты на запас, ₽', 'Полная оценка, ₽', 'Оценено бутылок', 'Известная часть оценки, ₽'])
            for record in history:
                history_sheet.append([safe_cell(record.get(column)) for column in history_columns])
            consumed_sheet = book.create_sheet('Выпито и результат')
            consumed_columns = ['name', 'winery', 'vintage', 'quantity', 'consumed_on', 'purchase_price',
                                'valuation_price', 'result', 'result_pct', 'price_date', 'price_source', 'price_url']
            consumed_sheet.append(['Вино', 'Производитель', 'Год урожая', 'Выпито, шт.', 'Дата списания',
                'Цена покупки за бутылку, ₽', 'Оценка при списании за бутылку, ₽', 'Зафиксированный условный результат, ₽',
                'Результат, %', 'Дата оценки', 'Источник оценки', 'Ссылка'])
            for record in cellar.consumptions(store, actor, owner=owner):
                consumed_sheet.append([safe_cell(record.get(column)) for column in consumed_columns])
            consumed_sheet.freeze_panes = 'A2'
            consumed_sheet.auto_filter.ref = consumed_sheet.dimensions
            for column in consumed_sheet.columns:
                consumed_sheet.column_dimensions[column[0].column_letter].width = min(55, max(16, max(len(str(cell.value or '')) for cell in column)+2))
            book.save(path)
        else:
            heads = ''.join('<th>'+html.escape(LABELS.get(column, column))+'</th>' for column in columns)
            body = ''.join('<tr>'+''.join('<td>'+html.escape(str(row.get(column) if row.get(column) is not None else '—'))+'</td>' for column in columns)+'</tr>' for row in rows)
            page = '''<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Моя винотека</title>
<style>body{font:16px system-ui;background:#f7f3ee;color:#362733;margin:0;padding:24px}h1{color:#713b52}p{max-width:900px;line-height:1.5}input{padding:12px;width:min(90%,500px);border:1px solid #cab9bd;border-radius:8px}.table{overflow:auto;margin-top:20px}table{border-collapse:collapse;background:white;min-width:900px}th,td{text-align:left;padding:12px;border-bottom:1px solid #eee;white-space:pre-wrap}th{background:#713b52;color:white;cursor:pointer;position:sticky;top:0}td{max-width:350px}small{color:#70616c}</style>
<h1>Моя винотека</h1><p>Личный снимок запасов. Цены относятся к указанной дате и источнику. Магазинная цена не означает стоимость обратной продажи. Прочерк означает отсутствие данных.</p>
<input id="search" placeholder="Найти вино, производителя или место хранения"><small> Нажмите заголовок для сортировки</small><div class="table"><table id="wines"><thead><tr>'''+heads+'</tr></thead><tbody>'+body+'''</tbody></table></div>
<script>const rows=[...document.querySelectorAll('tbody tr')];document.getElementById('search').oninput=e=>{const q=e.target.value.toLocaleLowerCase();rows.forEach(r=>r.hidden=!r.textContent.toLocaleLowerCase().includes(q));};document.querySelectorAll('th').forEach((h,i)=>h.onclick=()=>{h.dataset.asc=h.dataset.asc==='1'?'0':'1';const sign=h.dataset.asc==='1'?1:-1;rows.sort((a,b)=>sign*a.cells[i].textContent.localeCompare(b.cells[i].textContent,'ru',{numeric:true})).forEach(r=>document.querySelector('tbody').append(r));});</script></html>'''
            path.write_text(page, encoding='utf-8')
        part = {'kind': 'document', 'path': str(path), 'caption': 'Моя винотека · '+format_.upper()}
    text = ('Моя винотека · '+('график' if format_ == 'chart' else 'таблица запасов')+
            f'\nПозиций: {len(rows)}. Снимок на {utcnow()[:10]}.\nПокупки и списания меняют размер запаса; '
            'изменение цен сравнивается отдельно для одних и тех же бутылок. Неизвестная стоимость не заменяется нулём.')
    if format_ == 'chart':
        text += '\nНа графике последнее сохранённое состояние каждого дня по Москве; все операции доступны в истории и Excel.'
    _result(store, actor, report_id, text, parts=[part], owner=owner)


def run_job(store, settings, job, payload, *, dry_run=False):
    actor = payload.get('actor_id')
    owner = cellar.require_owner(store, actor, payload.get('owner_id'))
    report_identifier(payload['report_id'])
    existing = store.one('SELECT quality FROM reports WHERE id=?', (payload['report_id'],))
    if existing:
        quality = json.loads(existing['quality'])
        if (quality.get('feature') != 'wine_cellar' or quality.get('actor_id') != actor
                or quality.get('owner_id') != owner):
            raise PermissionError('Этот результат принадлежит другому пользователю или инструменту.')
    else:
        if job['kind'] == 'wine_recognize':
            _recognize(store, settings, job, payload, dry_run=dry_run)
        elif job['kind'] == 'wine_price_refresh':
            _prices(store, job, payload, dry_run=dry_run)
        elif job['kind'] == 'wine_price_search':
            _search(store, payload, dry_run=dry_run)
        elif job['kind'] == 'wine_cellar_export':
            _export(store, settings, payload)
        elif job['kind'] == 'wine_details_lookup':
            from .wine_details import lookup
            result = lookup(store,actor,payload['lot_id'],payload['details_token'],dry_run=dry_run)
            text = ('Поиск сведений для сохранённого вина завершён. Карточек: '+str(len(result.get('candidates',[])))+
                    '. Откройте уточнение, выберите нужное поле и проверьте варианты по этикетке. '
                    'Изменения сохраняются отдельной кнопкой; покупка и запасы не добавляются повторно.')
            _result(store,actor,payload['report_id'],text,buttons=[[
                {'text':'Уточнить характеристики','callback_data':f"wc:details:{payload['lot_id']}"}]])
        else:
            raise ValueError('Неизвестное задание винотеки')
    if not dry_run:
        prepare_delivery(store, payload['report_id'], str(actor), detail='full')
    return payload['report_id']
